import csv
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from fastapi import HTTPException, status
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.models.business import InventoryTxn, Reagent, ReagentLot
from app.models.user import User
from app.services.reagents import ensure_can_manage_reagent_master


IMPORT_HEADERS = [
    "试剂名称",
    "批号",
    "库存数量",
    "单位",
    "试剂编码",
    "CAS号",
    "供应商",
    "货号",
    "规格",
    "存放位置",
    "储存条件",
    "入库日期",
    "有效期至",
    "负责人",
    "备注",
]
REQUIRED_HEADERS = {"试剂名称", "批号", "库存数量", "单位"}
MAX_FILE_SIZE = 2 * 1024 * 1024
MAX_ROWS = 1000


@dataclass
class ParsedRow:
    row_number: int
    values: dict[str, str]
    quantity: Decimal
    received_date: date | None
    expiry_date: date | None
    reagent: Reagent | None
    reagent_key: tuple[Any, ...]


def build_template(template_format: str) -> tuple[bytes, str, str]:
    if template_format == "csv":
        stream = StringIO()
        csv.writer(stream).writerow(IMPORT_HEADERS)
        return (
            stream.getvalue().encode("utf-8-sig"),
            "text/csv; charset=utf-8",
            "reagent-import-template.csv",
        )
    if template_format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "试剂库存导入"
        sheet.append(IMPORT_HEADERS)
        stream = BytesIO()
        workbook.save(stream)
        return (
            stream.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "reagent-import-template.xlsx",
        )
    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="模板格式仅支持 csv 或 xlsx")


def import_reagents(
    db: Session,
    current_user: User,
    *,
    filename: str | None,
    content: bytes,
    dry_run: bool,
) -> dict:
    ensure_can_manage_reagent_master(current_user)
    extension = Path(filename or "").suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="仅支持 .csv 和 .xlsx 文件")
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="文件大小不能超过 2MB")

    headers, raw_rows = _read_rows(extension, content)
    normalized_headers = [_normalize_header(value) for value in headers]
    missing_headers = [header for header in IMPORT_HEADERS if header not in normalized_headers]
    if missing_headers:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"模板表头不完整，缺少：{'、'.join(missing_headers)}",
        )

    rows = _map_rows(normalized_headers, raw_rows)
    if len(rows) > MAX_ROWS:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="导入数据不能超过 1000 行")

    result, parsed_rows = _validate_rows(db, rows, dry_run)
    if dry_run or result["errors"]:
        if not dry_run:
            result["created_reagents"] = 0
            result["matched_reagents"] = 0
            result["created_lots"] = 0
        return result

    _persist_rows(db, current_user, parsed_rows, result)
    return result


def _read_rows(extension: str, content: bytes) -> tuple[list[Any], list[list[Any]]]:
    try:
        if extension == ".csv":
            text = content.decode("utf-8-sig")
            rows = list(csv.reader(StringIO(text)))
        else:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            rows = [list(row) for row in workbook.active.iter_rows(values_only=True)]
    except (UnicodeDecodeError, csv.Error, ValueError, OSError, KeyError, BadZipFile) as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="文件内容无法解析") from exc
    if not rows:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="文件缺少表头")
    return rows[0], rows[1:]


def _normalize_header(value: Any) -> str:
    return str(value or "").lstrip("\ufeff").strip()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _map_rows(headers: list[str], raw_rows: list[list[Any]]) -> list[tuple[int, dict[str, str]]]:
    mapped_rows: list[tuple[int, dict[str, str]]] = []
    for row_number, raw_row in enumerate(raw_rows, start=2):
        values = {
            header: _cell_text(raw_row[index] if index < len(raw_row) else None)
            for index, header in enumerate(headers)
        }
        selected = {header: values.get(header, "") for header in IMPORT_HEADERS}
        if any(selected.values()):
            mapped_rows.append((row_number, selected))
    return mapped_rows


def _error(row: int, field: str, message: str) -> dict:
    return {"row": row, "field": field, "message": message}


def _warning(row: int, field: str, message: str) -> dict:
    return {"row": row, "field": field, "message": message}


def _parse_date(row_number: int, field: str, value: str, errors: list[dict]) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError:
        errors.append(_error(row_number, field, f"{field}必须是合法日期，格式为 YYYY-MM-DD"))
        return None


def _find_existing_reagent(db: Session, name: str, cas_no: str) -> Reagent | None:
    stmt = select(Reagent).where(Reagent.name == name)
    if cas_no:
        stmt = stmt.where(Reagent.cas_no == cas_no)
    return db.scalar(stmt.order_by(Reagent.id).limit(1))


def _validate_rows(
    db: Session,
    rows: list[tuple[int, dict[str, str]]],
    dry_run: bool,
) -> tuple[dict, list[ParsedRow]]:
    errors: list[dict] = []
    warnings: list[dict] = []
    parsed_rows: list[ParsedRow] = []
    duplicate_keys: set[tuple[tuple[Any, ...], str]] = set()
    new_reagent_keys: set[tuple[Any, ...]] = set()
    matched_rows = 0
    error_row_numbers: set[int] = set()

    for row_number, values in rows:
        row_errors: list[dict] = []
        for field in REQUIRED_HEADERS:
            if not values[field]:
                row_errors.append(_error(row_number, field, f"{field}为必填字段"))

        try:
            quantity = Decimal(values["库存数量"])
            if not quantity.is_finite() or quantity < 0:
                raise InvalidOperation
        except (InvalidOperation, ValueError):
            quantity = Decimal("0")
            if values["库存数量"]:
                row_errors.append(_error(row_number, "库存数量", "库存数量必须是大于等于 0 的数字"))

        received_date = _parse_date(row_number, "入库日期", values["入库日期"], row_errors)
        expiry_date = _parse_date(row_number, "有效期至", values["有效期至"], row_errors)
        reagent = None
        name = values["试剂名称"]
        cas_no = values["CAS号"]
        if name:
            reagent = _find_existing_reagent(db, name, cas_no)

        if values["试剂编码"]:
            warnings.append(
                _warning(row_number, "试剂编码", "现有模型无试剂编码字段，无法按编码匹配或保存，已按试剂名称和CAS号匹配")
            )
        if not cas_no:
            warnings.append(_warning(row_number, "CAS号", "CAS号为空，已按试剂名称尝试匹配"))

        reagent_key = ("existing", reagent.id) if reagent else ("new", name, cas_no)
        lot_key = (reagent_key, values["批号"])
        if name and values["批号"]:
            if lot_key in duplicate_keys:
                row_errors.append(_error(row_number, "批号", "文件内重复：同一试剂的批号重复"))
            else:
                duplicate_keys.add(lot_key)
            if reagent is not None and db.scalar(
                select(ReagentLot.id)
                .where(ReagentLot.reagent_id == reagent.id, ReagentLot.lot_no == values["批号"])
                .limit(1)
            ):
                row_errors.append(_error(row_number, "批号", "同一试剂下该批号已存在"))

        for field in ("入库日期", "负责人", "备注"):
            if values[field]:
                warnings.append(_warning(row_number, field, f"现有模型无{field}字段，该值未保存"))

        if row_errors:
            errors.extend(row_errors)
            error_row_numbers.add(row_number)
            continue

        if reagent is not None:
            matched_rows += 1
        else:
            new_reagent_keys.add(reagent_key)
        parsed_rows.append(
            ParsedRow(
                row_number=row_number,
                values=values,
                quantity=quantity,
                received_date=received_date,
                expiry_date=expiry_date,
                reagent=reagent,
                reagent_key=reagent_key,
            )
        )

    valid_rows = len(rows) - len(error_row_numbers)
    result = {
        "dry_run": dry_run,
        "total_rows": len(rows),
        "valid_rows": valid_rows,
        "error_rows": len(error_row_numbers),
        "created_reagents": len(new_reagent_keys),
        "matched_reagents": matched_rows,
        "created_lots": valid_rows,
        "errors": errors,
        "warnings": warnings,
    }
    return result, parsed_rows


def _persist_rows(db: Session, current_user: User, rows: list[ParsedRow], result: dict) -> None:
    created_reagents: dict[tuple[Any, ...], Reagent] = {}
    try:
        for parsed in rows:
            reagent = parsed.reagent or created_reagents.get(parsed.reagent_key)
            if reagent is None:
                reagent = Reagent(
                    name=parsed.values["试剂名称"],
                    cas_no=parsed.values["CAS号"] or None,
                    catalog_no=parsed.values["货号"] or None,
                    manufacturer=parsed.values["供应商"] or None,
                    grade=parsed.values["规格"] or None,
                    default_unit=parsed.values["单位"],
                    is_active=True,
                    created_by=current_user.id,
                )
                db.add(reagent)
                db.flush()
                created_reagents[parsed.reagent_key] = reagent

            lot = ReagentLot(
                reagent_id=reagent.id,
                lot_no=parsed.values["批号"],
                expiry_date=parsed.expiry_date,
                quantity=parsed.quantity,
                unit=parsed.values["单位"],
                location=parsed.values["存放位置"] or None,
                storage_condition=parsed.values["储存条件"] or None,
                status="in_stock",
                created_by=current_user.id,
            )
            db.add(lot)
            db.flush()
            if parsed.quantity > 0:
                db.add(
                    InventoryTxn(
                        reagent_lot_id=lot.id,
                        txn_type="in",
                        quantity=parsed.quantity,
                        balance_after=parsed.quantity,
                        reference="试剂库存导入",
                        operator_id=current_user.id,
                        source_type="import",
                    )
                )
        db.commit()
    except SQLAlchemyError as exc:
        db.rollback()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="导入失败，所有数据已回滚") from exc
