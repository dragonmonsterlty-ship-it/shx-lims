import csv
from io import BytesIO, StringIO

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.models.business import Reagent, ReagentLot

from test_reagents import auth_headers, create_lot, create_reagent, setup_users


HEADERS = [
    "试剂名称",
    "批号",
    "库存数量",
    "单位",
    "试剂编码",
    "CAS号",
    "供应商",
    "货号",
    "规格",
    "存放位置",
    "储存条件",
    "入库日期",
    "有效期至",
    "负责人",
    "备注",
]


def csv_file(rows: list[list[object]], headers: list[str] = HEADERS) -> tuple[str, bytes, str]:
    stream = StringIO()
    writer = csv.writer(stream)
    writer.writerow(headers)
    writer.writerows(rows)
    return ("reagents.csv", stream.getvalue().encode("utf-8-sig"), "text/csv")


def xlsx_file(rows: list[list[object]], headers: list[str] = HEADERS) -> tuple[str, bytes, str]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return (
        "reagents.xlsx",
        stream.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def import_file(client, headers, upload, dry_run=True):
    return client.post(
        "/api/reagents/import",
        params={"dry_run": str(dry_run).lower()},
        headers=headers,
        files={"file": upload},
    )


def row(
    name="甲醇",
    lot_no="LOT-001",
    quantity="10.5",
    unit="mL",
    reagent_code="",
    cas_no="67-56-1",
    supplier="供应商A",
    catalog_no="CAT-001",
    specification="HPLC",
    location="A柜",
    storage_condition="室温",
    received_date="2026-06-01",
    expiry_date="2027-06-01",
    owner="张三",
    remark="避光",
):
    return [
        name,
        lot_no,
        quantity,
        unit,
        reagent_code,
        cas_no,
        supplier,
        catalog_no,
        specification,
        location,
        storage_condition,
        received_date,
        expiry_date,
        owner,
        remark,
    ]


def test_csv_template_has_bom_and_exact_chinese_headers(client, create_user):
    setup_users(create_user)

    response = client.get(
        "/api/reagents/import-template",
        params={"format": "csv"},
        headers=auth_headers(client, "admin"),
    )

    assert response.status_code == 200
    assert response.content.startswith(b"\xef\xbb\xbf")
    assert next(csv.reader(StringIO(response.content.decode("utf-8-sig")))) == HEADERS


def test_xlsx_template_has_exact_chinese_headers(client, create_user):
    setup_users(create_user)

    response = client.get(
        "/api/reagents/import-template",
        params={"format": "xlsx"},
        headers=auth_headers(client, "director"),
    )

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert list(next(workbook.active.iter_rows(values_only=True))) == HEADERS


def test_csv_dry_run_validates_without_writing(client, create_user, db_session):
    setup_users(create_user)

    response = import_file(client, auth_headers(client, "admin"), csv_file([row()]))

    assert response.status_code == 200
    data = response.json()["data"]
    assert data["dry_run"] is True
    assert data["total_rows"] == 1
    assert data["valid_rows"] == 1
    assert data["error_rows"] == 0
    assert data["created_reagents"] == 1
    assert data["matched_reagents"] == 0
    assert data["created_lots"] == 1
    assert db_session.scalar(select(func.count()).select_from(Reagent)) == 0
    assert db_session.scalar(select(func.count()).select_from(ReagentLot)) == 0


def test_xlsx_dry_run_skips_empty_rows_and_accepts_dates(client, create_user):
    setup_users(create_user)

    response = import_file(client, auth_headers(client, "admin"), xlsx_file([row(), [""] * len(HEADERS)]))

    assert response.status_code == 200
    assert response.json()["data"]["total_rows"] == 1
    assert response.json()["data"]["valid_rows"] == 1


def test_formal_import_creates_reagent_and_lot_with_supported_fields(client, create_user, db_session):
    setup_users(create_user)

    response = import_file(client, auth_headers(client, "director"), csv_file([row()]), dry_run=False)

    assert response.status_code == 200
    data = response.json()["data"]
    assert data["created_reagents"] == 1
    assert data["created_lots"] == 1
    db_session.expire_all()
    reagent = db_session.scalar(select(Reagent))
    lot = db_session.scalar(select(ReagentLot))
    assert (reagent.name, reagent.cas_no, reagent.manufacturer, reagent.catalog_no, reagent.grade) == (
        "甲醇",
        "67-56-1",
        "供应商A",
        "CAT-001",
        "HPLC",
    )
    assert (lot.lot_no, str(lot.quantity), lot.unit, lot.location, lot.storage_condition) == (
        "LOT-001",
        "10.5000",
        "mL",
        "A柜",
        "室温",
    )
    assert lot.expiry_date.isoformat() == "2027-06-01"
    warning_fields = {warning["field"] for warning in data["warnings"]}
    assert {"入库日期", "负责人", "备注"} <= warning_fields


def test_missing_required_field_returns_row_error(client, create_user):
    setup_users(create_user)

    response = import_file(client, auth_headers(client, "admin"), csv_file([row(unit="")]))

    assert response.status_code == 200
    assert response.json()["data"]["errors"] == [{"row": 2, "field": "单位", "message": "单位为必填字段"}]


def test_non_numeric_and_negative_quantity_return_errors(client, create_user):
    setup_users(create_user)
    headers = auth_headers(client, "admin")

    non_numeric = import_file(client, headers, csv_file([row(quantity="abc")]))
    negative = import_file(client, headers, csv_file([row(quantity="-1")]))

    expected = "库存数量必须是大于等于 0 的数字"
    assert non_numeric.json()["data"]["errors"][0]["message"] == expected
    assert negative.json()["data"]["errors"][0]["message"] == expected


def test_invalid_dates_return_errors(client, create_user):
    setup_users(create_user)

    response = import_file(
        client,
        auth_headers(client, "admin"),
        csv_file([row(received_date="2026-02-30", expiry_date="not-a-date")]),
    )

    assert [(item["field"], item["message"]) for item in response.json()["data"]["errors"]] == [
        ("入库日期", "入库日期必须是合法日期，格式为 YYYY-MM-DD"),
        ("有效期至", "有效期至必须是合法日期，格式为 YYYY-MM-DD"),
    ]


def test_database_duplicate_lot_is_rejected(client, create_user):
    setup_users(create_user)
    headers = auth_headers(client, "admin")
    reagent = create_reagent(client, headers, name="甲醇")
    create_lot(client, headers, reagent["id"], "LOT-001")

    response = import_file(client, headers, csv_file([row()]))

    assert response.json()["data"]["errors"][0]["field"] == "批号"
    assert "已存在" in response.json()["data"]["errors"][0]["message"]


def test_file_internal_duplicate_lot_is_rejected(client, create_user):
    setup_users(create_user)

    response = import_file(
        client,
        auth_headers(client, "admin"),
        csv_file([row(), row(supplier="另一供应商")]),
    )

    assert response.json()["data"]["error_rows"] == 1
    assert response.json()["data"]["errors"][0]["row"] == 3
    assert "文件内重复" in response.json()["data"]["errors"][0]["message"]


def test_invalid_file_type_is_rejected(client, create_user):
    setup_users(create_user)

    response = import_file(
        client,
        auth_headers(client, "admin"),
        ("reagents.xls", b"not-an-xls", "application/vnd.ms-excel"),
    )

    assert response.status_code == 400
    assert response.json()["message"] == "仅支持 .csv 和 .xlsx 文件"


def test_corrupt_xlsx_returns_safe_parse_error(client, create_user):
    setup_users(create_user)

    response = import_file(
        client,
        auth_headers(client, "admin"),
        (
            "reagents.xlsx",
            b"not-a-valid-workbook",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
    )

    assert response.status_code == 400
    assert response.json()["message"] == "文件内容无法解析"


def test_operator_and_manager_can_import_while_unauthenticated_is_forbidden(client, create_user):
    setup_users(create_user)
    upload = csv_file([row()])

    unauthenticated = import_file(client, {}, upload)
    operator_template = client.get(
        "/api/reagents/import-template",
        params={"format": "csv"},
        headers=auth_headers(client, "operator"),
    )
    operator_import = import_file(client, auth_headers(client, "operator"), upload)
    manager_import = import_file(
        client,
        auth_headers(client, "manager"),
        csv_file([row(lot_no="LOT-MANAGER")]),
    )

    assert unauthenticated.status_code == 401
    assert operator_template.status_code == 200
    assert operator_import.status_code == 200
    assert operator_import.json()["data"]["valid_rows"] == 1
    assert manager_import.status_code == 200
    assert manager_import.json()["data"]["valid_rows"] == 1


def test_formal_import_with_any_error_rolls_back_all_rows(client, create_user, db_session):
    setup_users(create_user)

    response = import_file(
        client,
        auth_headers(client, "admin"),
        csv_file([row(), row(name="乙醇", lot_no="LOT-002", quantity="-1", cas_no="64-17-5")]),
        dry_run=False,
    )

    assert response.status_code == 200
    assert response.json()["data"]["error_rows"] == 1
    assert response.json()["data"]["created_lots"] == 0
    assert db_session.scalar(select(func.count()).select_from(Reagent)) == 0
    assert db_session.scalar(select(func.count()).select_from(ReagentLot)) == 0


def test_trimmed_bom_header_and_existing_reagent_match(client, create_user):
    setup_users(create_user)
    headers = auth_headers(client, "admin")
    create_reagent(client, headers, name="甲醇")
    padded_headers = ["\ufeff 试剂名称 "] + [f" {header} " for header in HEADERS[1:]]

    response = import_file(client, headers, csv_file([row()], padded_headers))

    data = response.json()["data"]
    assert data["valid_rows"] == 1
    assert data["created_reagents"] == 0
    assert data["matched_reagents"] == 1


def test_missing_template_header_is_rejected(client, create_user):
    setup_users(create_user)

    response = import_file(client, auth_headers(client, "admin"), csv_file([row()[:-1]], HEADERS[:-1]))

    assert response.status_code == 400
    assert response.json()["message"].startswith("模板表头不完整")
